"""Generate the ClosureCopilot pitch deck (PPTX)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PRIMARY = RGBColor(0x1F, 0x6F, 0xEB)
PURPLE = RGBColor(0x89, 0x57, 0xE5)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x57, 0x60, 0x6A)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2D, 0xA4, 0x4E)
RED = RGBColor(0xE5, 0x48, 0x4D)
ORANGE = RGBColor(0xC8, 0x86, 0x0A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, (t, size, color, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
        r.font.name = "Segoe UI"
    return tb


def bg(s, color=WHITE):
    rect(s, 0, 0, SW, SH, color)


def header(s, title, accent=PRIMARY):
    rect(s, 0, 0, SW, Inches(1.15), accent)
    text(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.8),
         [(title, 30, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)


# ---- Slide 1 : Title ----
s = slide(); bg(s)
rect(s, 0, 0, SW, SH, RGBColor(0x0F, 0x2B, 0x5C))
rect(s, 0, Inches(4.9), SW, Inches(0.12), PRIMARY)
text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.4),
     [("ClosureCopilot", 60, WHITE, True)])
text(s, Inches(0.95), Inches(3.35), Inches(11.5), Inches(0.8),
     [("A Multi-Agent Backend Closure Copilot for RTL", 26, RGBColor(0xBF, 0xD4, 0xF2), True)])
text(s, Inches(0.95), Inches(4.2), Inches(11.5), Inches(0.8),
     [("It decides WHERE each fix belongs — RTL  ·  SDC  ·  UPF", 20, RGBColor(0x9F, 0xB8, 0xE0), False)])

# ---- Slide 2 : Problem ----
s = slide(); bg(s); header(s, "The problem: PPA closure doesn't scale", RED)
bullets = [
    "Power, Performance & Area (PPA) closure is the biggest schedule risk in silicon.",
    "Run backend tools → wait hours → face thousands of lines of reports.",
    "The bottleneck is judgment, not reading:  is a timing violation a real RTL bug, or a missing constraint?  Is high clock power a hotspot, or unrealized UPF intent?",
    "That expertise lives only in senior engineers' heads — and doesn't scale.",
]
y = 1.7
for b in bullets:
    rect(s, Inches(0.7), Inches(y + 0.08), Inches(0.16), Inches(0.16), RED)
    text(s, Inches(1.05), Inches(y - 0.12), Inches(11.4), Inches(1.1),
         [(b, 19, DARK, False)])
    y += 1.15

# ---- Slide 3 : What it does (architecture) ----
s = slide(); bg(s); header(s, "One copilot · 9 agents · any backend report")
try:
    s.shapes.add_picture("docs/architecture.png", Inches(1.2), Inches(1.35),
                         width=Inches(11.0))
except Exception as e:
    text(s, Inches(1), Inches(3), Inches(11), Inches(1), [(f"[diagram] {e}", 14, GRAY, False)])
text(s, Inches(0.7), Inches(6.95), Inches(12), Inches(0.4),
     [("Tool-agnostic parsers normalize STA · synthesis · power · area · UPF · SDC · congestion into one finding model.",
       13, GRAY, False)], align=PP_ALIGN.CENTER)

# ---- Slide 4 : Fix-layer routing ----
s = slide(); bg(s); header(s, "The core idea: Fix-Layer Routing", PURPLE)
text(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.7),
     [("For every issue, it decides WHERE the fix belongs — RTL · SDC · UPF · synthesis-setup.",
       20, DARK, True)])
# example cards
rect(s, Inches(0.7), Inches(2.4), Inches(5.8), Inches(2.4), LIGHT)
text(s, Inches(0.95), Inches(2.6), Inches(5.3), Inches(2.1),
     [("Example 1 — fix in SDC (not RTL)", 17, PRIMARY, True),
      ("Config path cfg_reg → status_reg fails setup.", 15, DARK, False),
      ("Diagnosis: quasi-static, sampled once per write.", 15, GRAY, False),
      ("→ set_multicycle_path -setup 2 …", 15, GREEN, True)])
rect(s, Inches(6.85), Inches(2.4), Inches(5.8), Inches(2.4), LIGHT)
text(s, Inches(7.1), Inches(2.6), Inches(5.3), Inches(2.1),
     [("Example 2 — fix in RTL", 17, PURPLE, True),
      ("Deep combinational cone at alu.sv:210.", 15, DARK, False),
      ("Diagnosis: genuine logic depth, single cycle.", 15, GRAY, False),
      ("→ pipeline the datapath  (timing ✓, area +2%)", 15, GREEN, True)])
text(s, Inches(0.7), Inches(5.2), Inches(12), Inches(0.8),
     [("Every fix ships with a corrected, copy-paste snippet + a cross-domain PPA trade-off.",
       18, DARK, True)])

# ---- Slide 5 : five modules ----
s = slide(); bg(s); header(s, "Five modules, one dashboard", GREEN)
mods = [
    ("PPA Analyzer", "Ranks issues and routes each fix to RTL / SDC / UPF with a snippet + trade-off.", PRIMARY),
    ("Constraint Promotion", "Reconciles IP → top constraints; generates corrected top-level SDC.", PURPLE),
    ("UPF Signoff", "Checks isolation, retention & clock-gating intent; emits corrected UPF.", ORANGE),
    ("Regression Detective", "Compares two runs; finds what regressed and the change behind it.", RED),
    ("Physical-Aware Feedback", "Turns P&R congestion into concrete RTL restructuring hints.", GREEN),
]
y = 1.5
for name, desc, col in mods:
    rect(s, Inches(0.7), Inches(y), Inches(0.18), Inches(0.95), col)
    rect(s, Inches(1.0), Inches(y), Inches(11.6), Inches(0.95), LIGHT)
    text(s, Inches(1.25), Inches(y + 0.08), Inches(11.2), Inches(0.85),
         [(name + "  —  ", 17, col, True)])
    text(s, Inches(1.25), Inches(y + 0.44), Inches(11.2), Inches(0.5),
         [(desc, 14.5, DARK, False)])
    y += 1.08

# ---- Slide 6 : how it works ----
s = slide(); bg(s); header(s, "Grounded — and always-on", PRIMARY)
rect(s, Inches(0.7), Inches(1.5), Inches(5.9), Inches(3.4), LIGHT)
text(s, Inches(0.95), Inches(1.7), Inches(5.4), Inches(3.1),
     [("Dual-RAG knowledge", 20, PRIMARY, True),
      ("① Global PPA knowledge", 16, DARK, True),
      ("Low-power / timing / area methodology, tool-report semantics, fix patterns.", 14, GRAY, False),
      ("② Design Context Memory", 16, DARK, True),
      ("A persistent memory of YOUR chip — specs, micro-arch, clocking, prior runs.", 14, GRAY, False),
      ("→ Feedback reflects your design intent, and improves each iteration.", 14, GREEN, True)])
rect(s, Inches(6.9), Inches(1.5), Inches(5.7), Inches(3.4), LIGHT)
text(s, Inches(7.15), Inches(1.7), Inches(5.2), Inches(3.1),
     [("Built to run anywhere", 20, PURPLE, True),
      ("Python · multi-agent orchestration", 16, DARK, False),
      ("Azure OpenAI for reasoning", 16, DARK, False),
      ("Streamlit product dashboard", 16, DARK, False),
      ("Deterministic OFFLINE mode", 16, DARK, True),
      ("Runs reliably with or without credentials — the demo never breaks.", 14, GRAY, False)])

# ---- Slide 7 : why unique + impact ----
s = slide(); bg(s); header(s, "Why it wins", PURPLE)
rect(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(2.0), RGBColor(0xEE, 0xE7, 0xFA))
text(s, Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.7),
     [("Unique", 20, PURPLE, True),
      ("It combines fix-layer routing + IP→top constraint promotion + run-to-run regression",
       17, DARK, False),
      ("attribution for the RTL owner — a combination no existing point tool brings together.",
       17, DARK, False)])
rect(s, Inches(0.7), Inches(3.8), Inches(11.9), Inches(2.6), RGBColor(0xE7, 0xF3, 0xE9))
text(s, Inches(1.0), Inches(4.0), Inches(11.3), Inches(2.3),
     [("Impact", 20, GREEN, True),
      ("Hours of manual, expert-only triage → minutes of ranked, grounded, copy-paste fixes.", 17, DARK, False),
      ("Puts senior-level closure judgment in every RTL designer's hands.", 17, DARK, False),
      ("Catches unrealized intent and lost constraints before the next costly backend run.", 17, DARK, False)])

# ---- Slide 8 : closing ----
s = slide()
rect(s, 0, 0, SW, SH, RGBColor(0x0F, 0x2B, 0x5C))
text(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.2),
     [("ClosureCopilot", 52, WHITE, True)])
text(s, Inches(0.95), Inches(3.9), Inches(11.5), Inches(0.8),
     [("Close faster. Fix smarter.", 26, RGBColor(0xBF, 0xD4, 0xF2), True)])
text(s, Inches(0.95), Inches(5.2), Inches(11.5), Inches(0.5),
     [("github.com/mayureshpiske19/ClosureCopilot", 18, RGBColor(0x9F, 0xB8, 0xE0), False)])

prs.save("docs/ClosureCopilot.pptx")
print("saved docs/ClosureCopilot.pptx", "slides:", len(prs.slides._sldIdLst))
